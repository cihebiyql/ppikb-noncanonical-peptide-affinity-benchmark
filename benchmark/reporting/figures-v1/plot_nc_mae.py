"""One fixed NC-core MAE figure, from explicitly root-reviewed metric summaries."""
# ============================================================
# CONFIGURATION — paths come from the explicit CLI, never discovery
# ============================================================
from pathlib import Path
OUTPUT_FORMATS = ('png','pdf','jpg','svg')
DPI = 300
EN_SERIF_FALLBACK = ('Times New Roman','Liberation Serif','DejaVu Serif','serif')
ERROR_BAR = 'sample SD, 10 training seeds'
FIGURE_INCHES = (12.4,4.8)
BAR_COLORS = ('#555555','#dddddd')
BAR_HATCHES = ('','///')
PRECISION = 2
MAX_INPUT_BYTES = 64*1024*1024
MAX_RECORDS = 500000
MAX_WALL_SECONDS = 180
METHODS = ('GraphDTA-GIN','MGraphDTA','DGDTA-CL','DeepNC-GEN','GS-DTA','RF-Morgan2-chiral-AAC-DP')
LABELS = ('GraphDTA\nGIN','MGraphDTA','DGDTA\nCL','DeepNC\nGEN','GS-DTA','RF\nClassic')
REGIMES = ('NC_only','standard_only')
PANELS = ('source_peptide_dependency_cold','target_reference_homology_cold')
N = dict(zip(PANELS,(303,332)))
HERE = Path(__file__).resolve().parent
# ============================================================
# END CONFIGURATION
# ============================================================
import argparse
import csv
import hashlib
import json
import math
import os
import platform
import sys
import threading

def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for x in iter(lambda:f.read(1024*1024),b''):h.update(x)
    return h.hexdigest()
def write(p,x):Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def finite_number(value,nonnegative=True):
    return type(value) in (int,float) and math.isfinite(value) and (not nonnegative or value>=0)

def select_rows(summary):
    if not isinstance(summary,list) or len(summary)>MAX_RECORDS:raise ValueError('Expected bounded reporter summary array')
    rows={};known=set(METHODS)|{'reference:train_mean','reference:train_median','reference:target_reference_train_median'}
    for r in summary:
        if not isinstance(r,dict):raise ValueError('Summary record must be object')
        if r.get('dataset_id') not in ('PPIKB-internal','PTM-external75','IgG-supplement13'):raise ValueError('Unknown summary dataset')
        if r.get('method_id') not in known:raise ValueError('Unknown method')
        if (r.get('dataset_id'),r.get('cohort'),r.get('aggregation'),r.get('metric'))!=('PPIKB-internal','NC_core','row_micro','MAE'):continue
        if r['method_id'].startswith('reference:'):continue
        panel,regime=r.get('panel'),r.get('regime')
        if panel not in PANELS or regime not in REGIMES:raise ValueError('Unknown selected panel/regime')
        category='auxiliary_handcrafted_RF' if r['method_id']==METHODS[-1] else 'five_prior_paper_NN'
        if r.get('category')!=category:raise ValueError('Wrong method category')
        key=(panel,r['method_id'],regime)
        if key in rows:raise ValueError('Duplicate selected aggregate')
        if type(r.get('completed_seed_count')) is not int or type(r.get('defined_seed_count')) is not int or r['completed_seed_count']!=10 or r['defined_seed_count']!=10:raise ValueError('All ten completed and defined seeds required')
        counts=r.get('n_by_seed')
        if not isinstance(counts,dict) or set(counts)!={str(i) for i in range(10)} or any(type(v) is not int or v!=N[panel] for v in counts.values()):raise ValueError('Fixed NC303/332 per-seed denominator required')
        if r.get('n_unit')!='observation_rows':raise ValueError('Observation-row units required')
        if r.get('missing_seeds')!=[]:raise ValueError('Missing-seed metadata must be empty')
        if not finite_number(r.get('mean')) or not finite_number(r.get('sample_sd')):raise ValueError('Finite nonnegative mean and sample SD required')
        if 'values_by_seed' in r and (set(r['values_by_seed'])!={str(i) for i in range(10)} or any(not finite_number(v) for v in r['values_by_seed'].values())):raise ValueError('Undefined or invalid individual MAE')
        rows[key]=dict(r)
    expected={(p,m,r) for p in PANELS for m in METHODS for r in REGIMES}
    if set(rows)!=expected:raise ValueError('Exactly 24 prescribed NC aggregates required')
    return [rows[(p,m,r)] for p in PANELS for m in METHODS for r in REGIMES]

def health(rows,input_bytes):
    columns=('category','dataset_id','panel','method_id','regime','cohort','aggregation','metric','mean','sample_sd','completed_seed_count','defined_seed_count','n_unit')
    lines=['# Data health — reviewed 24 aggregate rows','',f'Rows: {len(rows)}; selected columns: {len(columns)}; full source bytes: {input_bytes}.',f'Selected serialized UTF-8 bytes (not in-memory RSS): {len(json.dumps(rows).encode())}.','', '| Column | Runtime types | Missing | Distinct |','| --- | --- | --- | --- |']
    for field in columns:
        vals=[r[field] for r in rows];lines.append(f"| {field} | {','.join(sorted({type(v).__name__ for v in vals}))} | {sum(v is None for v in vals)} | {len({json.dumps(v,sort_keys=True) for v in vals})} |")
    lines+=['','Duplicate panel/method/regime coordinates: 0. All 24 required groups present.','Each group has completed=defined=10; n_by_seed is 303 or 332 in the matching panel. Mean and sample SD are finite/nonnegative, retained directly without recalculation.','',f"MAE mean range: {min(r['mean'] for r in rows)!r} to {max(r['mean'] for r in rows)!r}; sample-SD range: {min(r['sample_sd'] for r in rows)!r} to {max(r['sample_sd'] for r in rows)!r}.",'','Normality tests, outlier deletion, numeric-predictor correlation and VIF: not applicable to this descriptive figure of already audited aggregates. No cleaning, fitting, seed selection, significance tests or summary recalculation.']
    return '\n'.join(lines)+'\n'

def export_data(rows,output):
    process=output/'process_data';process.mkdir()
    data=[]
    for r in rows:
        x={k:r[k] for k in ('category','dataset_id','panel','method_id','regime','cohort','aggregation','metric','mean','sample_sd','completed_seed_count','defined_seed_count','n_unit')}
        x.update(n=N[r['panel']],n_by_seed=json.dumps(r['n_by_seed'],sort_keys=True),mean_original_json=json.dumps(r['mean']),sample_sd_original_json=json.dumps(r['sample_sd']))
        data.append(x)
    columns=list(data[0])
    with (process/'fig_01_data.csv').open('w',encoding='utf-8-sig',newline='') as h:
        w=csv.DictWriter(h,fieldnames=columns);w.writeheader();w.writerows(data)
    write(process/'fig_01_original_aggregate_rows.json',rows)
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:return {'available':False,'reason':'openpyxl absent; no installation attempted'}
    book=openpyxl.Workbook();sheet=book.active;sheet.title='fig_01_data';sheet.append(columns)
    for x in data:sheet.append([x[k] for k in columns])
    sheet.freeze_panes='A2'
    for i,k in enumerate(columns,1):sheet.column_dimensions[get_column_letter(i)].width=min(50,max(len(str(k)),max(len(str(x[k])) for x in data))+2)
    book.save(process/'fig_01_data.xlsx')
    return {'available':True,'version':openpyxl.__version__,'exact_value_text_columns':'mean_original_json/sample_sd_original_json retain exact JSON numeric spellings alongside numeric Excel cells'}

def draw(rows,output,synthetic=False):
    os.environ.setdefault('MPLBACKEND','Agg')
    import matplotlib
    matplotlib.use('Agg',force=True)
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.ticker import FormatStrFormatter
    available={f.name for f in font_manager.fontManager.ttflist}
    font=next(x for x in EN_SERIF_FALLBACK if x in available or x=='serif')
    plt.rcParams.update({'font.family':'serif','font.serif':[font,'DejaVu Serif'],'font.size':10,'axes.labelsize':11,'xtick.labelsize':9,'ytick.labelsize':9,'axes.unicode_minus':False,'figure.facecolor':'white','axes.facecolor':'white','legend.frameon':False,'svg.fonttype':'none','pdf.fonttype':42,'axes.axisbelow':True})
    fig,axes=plt.subplots(1,2,figsize=FIGURE_INCHES,sharey=True)
    lookup={(r['panel'],r['method_id'],r['regime']):r for r in rows}
    ymin=min(0,min(r['mean']-r['sample_sd'] for r in rows));ymax=max(r['mean']+r['sample_sd'] for r in rows)
    span=max(ymax-ymin,0.1)
    for i,(ax,panel) in enumerate(zip(axes,PANELS)):
        for j,regime in enumerate(REGIMES):
            rs=[lookup[(panel,m,regime)] for m in METHODS];x=[n+(-.19 if j==0 else .19) for n in range(6)]
            ax.bar(x,[r['mean'] for r in rs],width=.34,yerr=[r['sample_sd'] for r in rs],label=regime,color=BAR_COLORS[j],hatch=BAR_HATCHES[j],edgecolor='black',linewidth=.65,capsize=3,error_kw={'linewidth':1,'ecolor':'black'})
        ax.set_xticks(range(6),LABELS);ax.set_xlabel('Fixed original-method order');ax.yaxis.set_major_formatter(FormatStrFormatter('%.'+str(PRECISION)+'f'));ax.yaxis.grid(True,color='#e5e5e5',linewidth=.5);ax.xaxis.grid(False)
        ax.set_ylim(ymin-(.04*span if ymin<0 else 0),ymax+.13*span)
        for side in ('top','right','left','bottom'):ax.spines[side].set_visible(True);ax.spines[side].set_linewidth(.8)
    axes[0].set_ylabel('MAE (pKd units) ↓')
    fig.subplots_adjust(left=.075,right=.985,bottom=.25,top=.75,wspace=.12)
    fig.text(.285,.79,'A  Source + peptide dependency cold; NC n=303',ha='center',fontsize=11)
    fig.text(.765,.79,'B  Target-reference homology cold; NC n=332',ha='center',fontsize=11)
    handles,labels=axes[0].get_legend_handles_labels();fig.legend(handles,labels,loc='upper center',bbox_to_anchor=(.53,.92),ncol=2,fontsize=10)
    fig.text(.53,.055,'Mean ± sample SD, 10 training seeds; SD reflects training randomness, not biological replication.',ha='center',fontsize=9)
    if synthetic:fig.text(.53,.985,'SYNTHETIC MECHANICAL TEST — NOT SCIENTIFIC RESULTS',ha='center',fontsize=10,color='black')
    for extension in OUTPUT_FORMATS:
        folder=output/'figures'/extension;folder.mkdir(parents=True)
        fig.savefig(folder/('fig_01_nc_mae.'+extension),dpi=DPI,bbox_inches='tight',facecolor='white')
    plt.close(fig)
    return {'matplotlib_version':matplotlib.__version__,'backend':matplotlib.get_backend(),'font':font,'dpi':DPI,'figure_inches':FIGURE_INCHES,'frame':'four_visible_spines','bar_colors':BAR_COLORS,'bar_hatches':BAR_HATCHES,'error_bars':ERROR_BAR,'shared_y_scale':True}

def produce(summary,output,input_sha,input_bytes,synthetic=False):
    rows=select_rows(summary);output=Path(output)
    output.mkdir(parents=True,exist_ok=False)
    (output/'00_data_health.md').write_text(health(rows,input_bytes),encoding='utf-8')
    xlsx=export_data(rows,output);render=draw(rows,output,synthetic)
    caption='图1．相同 NC_core 测试集上的亲和力误差比较。左侧为来源＋肽依赖冷面板（303条），右侧为参考靶标同源冷面板（332条）；每侧按固定作者顺序比较五个NN架构与一个普通RF基准的NC_only及standard_only训练结果。柱为10个训练随机种子的MAE均值，误差条为样本标准差，均直接取自root审核的汇总表；MAE越低越好。SD描述训练随机性，两个面板分别解释，10个种子不是独立生物重复；图中不使用all测试池或确定性参考规则替代NC主结果。'
    if synthetic:caption='【纯合成机械测试，无科学成绩】\n'+caption
    (output/'caption-zh.md').write_text(caption+'\n',encoding='utf-8')
    (output/'README.md').write_text(('纯合成机械测试。\n' if synthetic else 'Root审核的真实汇总指标图。\n')+'\nfigures/下为300dpi PNG/JPG、矢量PDF/SVG；中文图注见caption-zh.md。process_data/保留24组精确图用CSV、原汇总JSON与可用的XLSX。00_data_health.md记录绘图前字段与分母验证。均值和样本SD直接使用原汇总，未选择seed或重算指标。\n',encoding='utf-8')
    receipt={'status':'synthetic_mechanical_only' if synthetic else 'complete_root_reviewed_aggregate_plot','summary_sha256':input_sha,'plotter_source_sha256':sha(Path(__file__)),'groups':24,'raw_gold_predictions_models_or_individual_training_labels_read':False,'summary_recalculated':False,'inferential_tests':False,'root_reviewed_aggregate_flag_required_for_actual':True,'python_version':platform.python_version(),'render':render,'xlsx':xlsx,'artifacts_sha256':{p.relative_to(output).as_posix():sha(p) for p in output.rglob('*') if p.is_file()}}
    write(output/'receipt.json',receipt)
    return receipt

def read_reviewed(args):
    if not args.root_reviewed_aggregate:raise PermissionError('Explicit --root-reviewed-aggregate required before summary access')
    source=Path(args.summary)
    if source.name not in ('summary.json','ten-seed-summary.json'):raise ValueError('Only reporter summary.json / ten-seed-summary.json input accepted')
    if source.stat().st_size>MAX_INPUT_BYTES:raise ValueError('Summary exceeds bounded input size')
    if not isinstance(args.summary_sha256,str) or len(args.summary_sha256)!=64 or sha(source)!=args.summary_sha256:raise ValueError('Reviewed summary SHA mismatch')
    if Path(args.output).resolve()==source.parent.resolve():raise ValueError('Separate figure output directory required')
    return json.loads(source.read_text(encoding='utf-8-sig')),source.stat().st_size

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root-reviewed-aggregate',action='store_true');parser.add_argument('--summary',required=True);parser.add_argument('--summary-sha256',required=True);parser.add_argument('--output',default=str(HERE/'actual-v1'));parser.add_argument('--wall-seconds',type=int,default=MAX_WALL_SECONDS)
    args=parser.parse_args()
    if not 1<=args.wall_seconds<=MAX_WALL_SECONDS:parser.error('Wall limit must be 1..180 seconds')
    def stop():
        sys.stderr.write('Plotter hard wall-clock limit reached; no continued process.\n');sys.stderr.flush();os._exit(124)
    timer=threading.Timer(args.wall_seconds,stop);timer.daemon=True;timer.start()
    try:
        summary,size=read_reviewed(args);receipt=produce(summary,args.output,args.summary_sha256,size)
        print(json.dumps({'status':receipt['status'],'groups':24,'output':args.output}))
    finally:timer.cancel()

if __name__=='__main__':main()

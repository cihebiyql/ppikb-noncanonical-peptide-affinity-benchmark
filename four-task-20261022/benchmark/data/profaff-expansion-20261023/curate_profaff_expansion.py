"""Bounded human-only, external-only ProfAff source curation. No model execution.

Run with the existing Python 3.11 + openpyxl environment. Sources are read only.
Composite holdup pKd is never relabelled as direct FP Kd. Zero sentinels are
censored observations, not zero affinities or confirmed non-interactions.
"""
from pathlib import Path
import collections
import csv
import datetime
import hashlib
import json
import math
import re
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
OLD = ROOT.parent / 'peptide-binding-benchmark-20261020/ten-seed-external-20261021/external/pepbi_pdz'
PPI = ROOT.parent / 'peptide-binding-expanded-corpus-20261021/resources/PPIKB_all_endpoints/normalized-observations.jsonl'
AA = set('ACDEFGHIKLMNPQRSTVWY')
DOI = '10.1038/s41467-022-33018-0'


def sha(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for block in iter(lambda: f.read(2**20), b''): h.update(block)
    return h.hexdigest()


def dump(name, x):
    (OUT / name).write_text(json.dumps(x, ensure_ascii=False, indent=2) + '\n', encoding='utf8')


def csvout(name, rows):
    assert rows, name
    with (OUT / name).open('w', encoding='utf8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)


def val(v):
    try:
        n = float(v)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError): return None


def norm(s): return s.upper().replace('-', '_').strip()


def chemistry(raw):
    """Only recognize actually observed, explicit tokens. No guessed chemistry."""
    s = raw.strip()
    m = re.fullmatch(r'(?i:biotin)-(ttds|ado-ado|Ahx|GSGSGS)-(.+)', s)
    if not m: return {'recognized': False, 'raw': raw}
    linker, core = m.groups()
    amidated = core.endswith('-CONH2')
    if amidated: core = core[:-6]
    # GSGSGS is itself a peptide linker and belongs to the supplied backbone.
    tokens = []; i = 0; ptms = []
    while i < len(core):
        found = next((t for t in ['Kac', 'pS', 'pT'] if core.startswith(t, i)), None)
        if found:
            tokens.append(found); ptms.append(found); i += len(found)
        elif core[i] in AA:
            tokens.append(core[i]); i += 1
        else: return {'recognized': False, 'raw': raw, 'reason': 'unrecognized_core_token'}
    parent = {'Kac': 'K', 'pS': 'S', 'pT': 'T'}
    backbone = ('GSGSGS' if linker == 'GSGSGS' else '') + ''.join(parent.get(t, t) for t in tokens)
    return {'recognized': True, 'core_tokens': tokens, 'core_parent_sequence': ''.join(parent.get(t,t) for t in tokens),
            'peptide_backbone': backbone, 'biotin_tag': True, 'linker': linker,
            'noncanonical_linker_Ahx': linker.lower() == 'ahx', 'noncanonical_core_residue': False,
            'ptm_tokens': ptms, 'terminal_amide': amidated,
            'projection_limitations': 'biotin/linker chemistry omitted from sequence input; PTM mapped to parent only in explicit projected track' if ptms else 'biotin/nonpeptide-linker chemistry and any terminal amidation omitted from sequence input',
            'chemical_complete_graph_available': False}


def main():
    sources = [OLD/'pdz-raw/41467_2022_33018_MOESM4_ESM.xlsx', OLD/'pdz-raw/database.sql',
               OLD/'pdz-raw/ProfAff-v1.0.zip', OLD/'pdz-human-raw-measurements.csv',
               OLD/'pdzome2023-domain-construct-sequences.json', OLD/'pdz-raw/pdzome2023-supplementary.pdf',
               OLD/'pdz-domain-source-records.json', OLD/'pdz.xml', OLD/'taxonomy-receipts.json', PPI,
               ROOT/'data/ppikb_kd_source_cold_scope_v2.csv', ROOT/'data/profaff_pKd_domain_backbone_all162_external.csv']
    sources += sorted((OLD/'taxonomy').glob('*.tsv')) + sorted((OUT/'sources').glob('uniprot-*.json'))
    source_hashes = [{'path': str(p), 'bytes': p.stat().st_size, 'sha256': sha(p)} for p in sources]
    before = {str(p): sha(p) for p in sources}
    human = set()
    for p in (OLD/'taxonomy').glob('*.tsv'):
        for x in csv.DictReader(p.open(encoding='utf8', newline=''), delimiter='\t'):
            if x.get('Organism (ID)') == '9606': human.add(x['Entry'])
    domains = {norm(x['target_source_name']): x for x in json.loads((OLD/'pdzome2023-domain-construct-sequences.json').read_text())}
    domain_meta = {x['PDZ']: x for x in json.loads((OLD/'pdz-domain-source-records.json').read_text())}
    aliases = {'AHNAK1':'AHNAK', 'NHERF1_1':'NHERF_1', 'NHERF1_2':'NHERF_2', 'TX1BP3':'TX1B3'}
    alias_audit = []
    for source_name, dest in aliases.items():
        acc = domain_meta[source_name]['UNIPROT']
        u = json.loads((OUT/'sources'/f'uniprot-{acc}.json').read_text())
        assert u['organism']['taxonId'] == 9606
        d = domains[dest]
        assert d['domain_sequence'] in u['sequence']['value']
        alias_audit.append({'source_name': source_name, 'source_accession': acc, 'construct_name': dest,
                            'uniprot_taxon': 9606, 'construct_exact_substring_of_same_accession': True,
                            'canonical_sequence_used_as_model_input': False, 'evidence_file': f'sources/uniprot-{acc}.json',
                            'does_not_verify_2022_fusion_construct': True})
    wb = openpyxl.load_workbook(sources[0], read_only=True, data_only=True)
    sheet = wb['PDZ-PBM interactome']
    header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=39, values_only=True))
    all_counts = collections.Counter(); records = []
    for n, values in enumerate(sheet.iter_rows(min_row=2, max_col=39, values_only=True), 2):
        r = {k: ('' if v is None else str(v)) for k,v in zip(header, values) if k}
        if not r.get('PDZ'): continue
        all_counts['all_contexts'] += 1
        if r['source']: all_counts['nonblank_source_contexts'] += 1
        if (val(r['COMP_neglogKd']) or 0) > 0: all_counts['positive_composite'] += 1
        if r['PBM_uniprot'].split('-')[0] not in human: continue
        r = {'source_excel_row': str(n), **r}
        records.append(r)
    wb.close()
    old_rows = {x['source_excel_row']:x for x in csv.DictReader((OLD/'pdz-human-raw-measurements.csv').open(encoding='utf8',newline=''))}
    crosscheck_errors = []
    for r in records:
        old = old_rows.get(r['source_excel_row'])
        if old is None: continue
        for k,v in r.items():
            if k not in old: continue
            a,b = val(v),val(old[k])
            equal = math.isclose(a,b,rel_tol=1e-8,abs_tol=1e-10) if a is not None and b is not None else v.strip()==old[k].strip()
            if not equal: crosscheck_errors.append([r['source_excel_row'],k,v,old[k]])
    assert not crosscheck_errors, crosscheck_errors[:5]
    assert len({(r['PDZ'],r['PBM']) for r in records}) == len(records)
    ppi = [json.loads(line) for line in PPI.open(encoding='utf8')]
    ref = list(csv.DictReader((ROOT/'data/ppikb_kd_source_cold_scope_v2.csv').open(encoding='utf8',newline='')))
    old_fp = list(csv.DictReader((ROOT/'data/profaff_pKd_domain_backbone_all162_external.csv').open(encoding='utf8',newline='')))
    # Keys from frozen evaluated FP rows are source row IDs; preserve as lineage.
    old_fp_ids = {re.search(r'(\d+)$',x['row_id']).group(1) for x in old_fp if re.search(r'(\d+)$',x['row_id'])}
    ref_pairs = {(x['target_sequence'],x['peptide_sequence']) for x in ref}
    ppi_seq = collections.defaultdict(list)
    for x in ppi:
        seq=x.get('peptide_sequence')
        if seq and set(seq)<=AA: ppi_seq[seq].append(x)
    ppi_targets = {x.get('target_sequence') for x in ppi if x.get('target_sequence')}
    ppi_studies = collections.Counter(str(x.get('study_id') or '').lower().replace('https://doi.org/','').strip() for x in ppi)
    points=[]; projected=[]; fp=[]; censored=[]; unknown=[]; sidecars=[]; overlap=[]
    chem_counts=collections.Counter(); target_map={}; matched_source_rows=0
    for r in records:
        dname=aliases.get(r['PDZ'],norm(r['PDZ'])); d=domains.get(dname)
        assert d is not None, r['PDZ']
        seq=d['domain_sequence']; assert set(seq)<=AA and len(seq)>=60
        target_map[r['PDZ']]={'construct_name':d['target_source_name'],'accession':domain_meta[r['PDZ']]['UNIPROT'],**d}
        c=chemistry(r['peptide']); assert c['recognized'],r['peptide']
        chem_counts['PTM_core' if c['ptm_tokens'] else 'standard_core']+=1
        chem_counts['linker_'+c['linker']]+=1
        if c['terminal_amide']: chem_counts['terminal_amide']+=1
        rid='ProfAff2022_HU_row'+r['source_excel_row']
        p=val(r['COMP_neglogKd']); f=val(r['FP_aver_uM']); measured=bool(r['source'])
        state='quantified_holdup' if measured and p is not None and p>0 else 'below_holdup_quantification' if measured and p==0 else 'not_measured_or_unresolved'
        exact=(seq,c['peptide_backbone']) in ref_pairs
        matches=[]
        for alt in {c['peptide_backbone'],r['peptide_10mer'],r['natural_10mer']}:
            if not alt or not set(alt)<=AA: continue
            for x in ppi_seq.get(alt,[]):
                ts=x.get('target_sequence') or ''
                if seq==ts or (ts and seq in ts):
                    matches.append({'PPIKB_record_id':x['record_id'],'study_id':x.get('study_id'),
                                    'sequence_relation':'exact_domain_input' if seq==ts else 'domain_contained_in_record_target',
                                    'matched_peptide_input':alt,'PPIKB_endpoint':x.get('endpoint'),
                                    'chemical_equivalence_verified':False})
        unique={json.dumps(m,sort_keys=True):m for m in matches}; matches=list(unique.values())
        if matches: overlap.append({'row_id':rid,'source_excel_row':r['source_excel_row'],'matches':matches})
        base={'row_id':rid,'source':'ProfAff2022_holdup','split':'external_test','peptide_sequence':c['peptide_backbone'],
              'peptide_sequence_raw':r['peptide'],'target_sequence':seq,'label':p if state=='quantified_holdup' else '',
              'endpoint':'holdup_derived_pKd','group_id':r['PDZ'],'peptide_smiles':'','smiles':'',
              'evaluation_stratum':'PTM_core_backbone_projection' if c['ptm_tokens'] else 'standard_core_assay_tagged_backbone_projection',
              'noncanonical_subtypes':';'.join(c['ptm_tokens']),'study_id':DOI,'target_id':domain_meta[r['PDZ']]['UNIPROT'],
              'chemical_entity_id':hashlib.sha256(r['peptide'].encode()).hexdigest(),'target_sequence_scope':'isolated_domain_2023_source_ancestry',
              'label_kind':'calibrated_holdup_composite_point' if state=='quantified_holdup' else state,
              'input_information_limitations':c['projection_limitations']+'; exact 2022 fusion sequence not verified',
              'source_excel_row':r['source_excel_row'],'target_name':r['PDZ'],'peptide_name':r['PBM'],
              'peptide_source_uniprot':r['PBM_uniprot'],'peptide_source_taxon':9606,'assay_method':r['COMP_method'],
              'source_codes':r['source'],'linker':c['linker'],'ptm_core':bool(c['ptm_tokens']),
              'noncanonical_core_residue':False,'noncanonical_linker_Ahx':c['noncanonical_linker_Ahx'],
              'biotin_tag':True,'terminal_amide':c['terminal_amide'],'actual_assay_construct_verified':False,
              'same_source_FP_available':bool(f is not None and f>0),
              'old_FP_162_source_row':r['source_excel_row'] in old_fp_ids,
              'exact_PPIKB_frozen_input_pair':exact,'PPIKB_possible_input_overlap':bool(matches),
              'new_blind_external_source':False,'chemically_complete_input':False}
        base['model_input_group_id'] = hashlib.sha256((seq+'|'+c['peptide_backbone']).encode()).hexdigest()
        base['chemical_measurement_group_id'] = hashlib.sha256((seq+'|'+r['peptide']).encode()).hexdigest()
        if state=='quantified_holdup':
            assert p>0 and math.isfinite(p)
            projected.append(base)
            if not c['ptm_tokens']: points.append(base)
        elif state=='below_holdup_quantification':
            # No universal numeric bound: PDZ concentration/calibration metadata
            # are not fully reconstructed here. Original per-assay fields remain.
            censored.append({**base,'label':'','pKd_relation':'below_assay_specific_quantification_limit',
                             'pKd_upper_bound':'','numeric_bound_status':'not_reconstructed',
                             'BI_detection_threshold':0.05 if r['COMP_method']=='DAPF' else 0.1 if r['COMP_method'] in ['LYSC','SAPF'] else '',
                             'confirmed_noninteraction':False})
        else: unknown.append({**base,'label':'','confirmed_noninteraction':False})
        if f is not None and f>0:
            fp.append({**base,'row_id':'ProfAff2022_FP_row'+r['source_excel_row'],'source':'ProfAff2022_directFP',
                       'label':6-math.log10(f),'endpoint':'direct_FP_pKd','label_kind':'direct_competitive_FP',
                       'Kd_uM':f,'Kd_SD_uM':r['FP_stdev_uM']})
        sidecars.append({'row_id':rid,'source_excel_row':r['source_excel_row'],'state':state,
                         'chemical_identity':r['peptide'],'chemistry':c,'raw_source_record':r,
                         'target_sequence_sha256':d['sequence_sha256'],'target_sequence_provenance':d,
                         'target_uniprot':domain_meta[r['PDZ']]['UNIPROT'],
                         'exact_2022_construct_equivalence_verified':False,
                         'calibration_warning':'Holdup uses effective peptide concentrations, partly FP calibrated and partly average concentration; shared calibration dependencies remain.',
                         'existing_162_FP_row':r['source_excel_row'] in old_fp_ids})
    assert len(projected)+len(censored)+len(unknown)==len(records)
    assert len({x['row_id'] for x in projected})==len(projected)
    assert all(x['label']>0 for x in projected)
    assert not any(x['label'] for x in censored+unknown)
    conservative = [x for x in points if not x['same_source_FP_available'] and not x['PPIKB_possible_input_overlap']]
    for name,rows in [('holdup-standard-core-external.csv',points),('holdup-standard-core-no-FP-no-PPIKB-overlap-external.csv',conservative),('holdup-all-explicit-backbone-projections-external.csv',projected),
                       ('direct-fp-separate-lineage.csv',fp),('holdup-censored-contexts.csv',censored),('unmeasured-or-unresolved-contexts.csv',unknown)]: csvout(name,rows)
    with (OUT/'source-lineage-chemistry.jsonl').open('w',encoding='utf8') as f:
        for x in sidecars:f.write(json.dumps(x,ensure_ascii=False)+'\n')
    dump('domain-input-provenance.json',list(target_map.values()))
    dump('alias-identity-audit.json',alias_audit)
    dump('PPIKB-possible-input-overlap.json',overlap)
    profile={'created_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'grain':'one original PDZ x PBM context; no combinatorial pairing',
             'original_workbook':dict(all_counts),'cached_taxon9606_accessions':len(human),'human_contexts':len(records),
             'old_human_contexts':len(old_rows),'new_human_rows_vs_old_filter':len(set(r['source_excel_row'] for r in records)-old_rows.keys()),
             'original_excel_vs_old_csv_crosschecked_rows':sum(r['source_excel_row'] in old_rows for r in records),
             'original_excel_vs_old_csv_errors':crosscheck_errors,'human_measured_contexts':len(projected)+len(censored),
             'human_quantified_holdup':len(projected),'human_censored_holdup':len(censored),'human_unmeasured_or_unresolved':len(unknown),
             'strict_standard_core_point_view':len(points),'explicit_PTM_parent_projection_point_rows':len(projected)-len(points),
             'conservative_standard_core_no_direct_FP_no_PPIKB_input_overlap':len(conservative),
             'human_direct_FP_positive':len(fp),'domain_inputs_resolved':len(target_map),'unique_chemical_strings':len({r['peptide'] for r in records}),
             'all_context_chemistry_counts':dict(chem_counts),'point_assay_counts':dict(collections.Counter(x['assay_method'] for x in projected)),
             'point_stratum_counts':dict(collections.Counter(x['evaluation_stratum'] for x in projected)),
             'point_linker_counts':dict(collections.Counter(x['linker'] for x in projected)),
             'point_unique_domains':len({x['target_name'] for x in projected}),'point_unique_peptide_names':len({x['peptide_name'] for x in projected}),
             'point_unique_chemical_strings':len({x['peptide_sequence_raw'] for x in projected}),
             'point_pKd_range':[min(x['label'] for x in projected),max(x['label'] for x in projected)],
             'point_same_source_FP':sum(x['same_source_FP_available'] for x in projected),
             'old_FP_id_parse_count':len(old_fp_ids),'PPIKB_all_endpoint_records':len(ppi),'PPIKB_same_publication_records':ppi_studies[DOI],
             'PPIKB_frozen_input_pair_matches_point':sum(x['exact_PPIKB_frozen_input_pair'] for x in projected),
             'PPIKB_possible_input_overlap_point':sum(x['PPIKB_possible_input_overlap'] for x in projected),
             'PPIKB_possible_input_overlap_all_contexts':len(overlap),
             'domains_exactly_contained_in_any_PPIKB_target':sum(any(d['domain_sequence'] in s for s in ppi_targets) for d in target_map.values()),
             'chemical_equivalence_overlap_not_proven':True,'full_2022_fusion_construct_verified':False,
             'blind_unseen_source':False,'new_scientific_forward_or_training_executed':False,
             'source_files_unchanged':all(sha(Path(p))==h for p,h in before.items())}
    assert profile['source_files_unchanged']
    dump('profile.json',profile);dump('source-hashes.json',source_hashes)
    outputs=[p for p in OUT.iterdir() if p.is_file() and p.name!='additional-manifest.json']
    dump('additional-manifest.json',{'schema':'ProfAff-human-external-expansion-v1','source_doi':DOI,'access':'external_only',
          'modifies_original_13_views':False,'purpose':'conditional external assay-shift evaluation; separate calibration-derived endpoint',
          'default_point_view':'holdup-standard-core-no-FP-no-PPIKB-overlap-external.csv',
          'broader_standard_core_point_view':'holdup-standard-core-external.csv','optional_lossy_PTM_projection_view':'holdup-all-explicit-backbone-projections-external.csv',
          'direct_FP_lineage_only':'direct-fp-separate-lineage.csv','never_pool_direct_FP_and_holdup_as_same_endpoint':True,
          'source_already_used_for_FP_evaluation':True,'blind_or_prospective_test':False,
          'numeric_censor_bounds_reconstructed':False,'all_input_graph_chemistry_available':False,
          'no_pairwise_effect_dataset_generated':True,'no_new_interaction_binary_or_contact_labels_inferred':True,
          'profile':'profile.json','source_hashes':'source-hashes.json',
          'files':[{'path':p.name,'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(outputs)]})
    print(json.dumps(profile,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
